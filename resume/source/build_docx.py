"""Build ATS-friendly .docx versions of the resume and cover letter.

Plain paragraph styles, real bullet lists, no tables or text boxes -- the
structures that resume parsers reliably handle.
"""
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT = "/tmp/claude-0/-home-user-General-Stuff/50909223-9957-552e-af9b-0843a2325d32/scratchpad/conv"


def new_doc(margin=0.6):
    doc = Document()
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Inches(margin)
        s.left_margin = s.right_margin = Inches(margin + 0.05)
    n = doc.styles["Normal"]
    n.font.name = "Calibri"
    n.font.size = Pt(10)
    n.paragraph_format.space_before = Pt(0)
    n.paragraph_format.space_after = Pt(0)
    n.paragraph_format.line_spacing = 1.0
    return doc


def runs(par, text):
    """Render **bold** segments as bold runs."""
    for i, chunk in enumerate(text.split("**")):
        if chunk:
            r = par.add_run(chunk)
            r.bold = i % 2 == 1


def name(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(18)
    return p


def contact(doc, text, rule=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8 if not rule else 10)
    r = p.add_run(text)
    r.font.size = Pt(9)
    if rule:
        bdr = OxmlElement("w:pBdr")
        bot = OxmlElement("w:bottom")
        bot.set(qn("w:val"), "single")
        bot.set(qn("w:sz"), "6")
        bot.set(qn("w:space"), "6")
        bot.set(qn("w:color"), "000000")
        bdr.append(bot)
        p._p.get_or_add_pPr().append(bdr)
    return p


def section(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(9)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text.upper())
    r.bold = True
    r.font.size = Pt(10.5)
    bdr = OxmlElement("w:pBdr")
    bot = OxmlElement("w:bottom")
    bot.set(qn("w:val"), "single")
    bot.set(qn("w:sz"), "6")
    bot.set(qn("w:space"), "2")
    bot.set(qn("w:color"), "000000")
    bdr.append(bot)
    p._p.get_or_add_pPr().append(bdr)
    return p


def job(doc, title, meta):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(7)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(title)
    r.bold = True
    r.font.size = Pt(10.5)
    m = doc.add_paragraph()
    m.paragraph_format.space_after = Pt(2)
    m.paragraph_format.keep_with_next = True
    mr = m.add_run(meta)
    mr.italic = True
    mr.font.size = Pt(9)


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.left_indent = Inches(0.22)
    p.paragraph_format.first_line_indent = Inches(-0.14)
    runs(p, text)


def body(doc, text, after=8):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    runs(p, text)
    return p


# --------------------------------------------------------------- resume

CONTACT = ("Columbus, Ohio (Remote)  |  mdgmikegeorgej@gmail.com  |  330-313-5604  |  "
           "linkedin.com/in/michael-george-02866411")

d = new_doc(0.5)
name(d, "MICHAEL GEORGE")
contact(d, CONTACT)

section(d, "Summary")
body(d, "Seventeen years in enterprise compute and virtualization. I currently run a 500+ "
        "server Windows and Linux estate across Hyper-V, VMware, and Nutanix for a national "
        "manufacturer, where I led the migration off VMware onto Hyper-V, built "
        "hyper-converged hosts on Storage Spaces Direct behind Windows Failover Clustering, "
        "and manage the environment day to day through Windows Admin Center and PowerShell. "
        "I've also delivered client-facing engagements as a consulting voice engineer, and "
        "spent six years as an IT Director owning infrastructure, budget, and vendor "
        "relationships outright. The documentation and runbooks I write are the ones teams "
        "actually open during an incident.", after=0)

section(d, "Experience")

job(d, "System Network Administrator — Overhead Door Corporation",
    "Columbus, OH  |  March 2019 – Present")
for b in [
    "**Led the migration from VMware to Hyper-V**, planning and executing the move of production workloads across multiple sites without interrupting the business.",
    "Built and operate **hyper-converged Hyper-V hosts on Storage Spaces Direct**, with Windows Failover Clustering behind the high-availability workloads.",
    "Manage **500+ Windows and Linux virtual servers** across multiple physical sites on Hyper-V with SCVMM, VMware vCenter and vSAN, and Nutanix.",
    "Standardized day-to-day host and server management on **Windows Admin Center**.",
    "Build and run virtual machines in **Azure** alongside the on-premises estate, and automate Azure and Oracle Cloud Infrastructure operations in **PowerShell**.",
    "Own **Active Directory** end to end — sites and services, replication topology, OU structure, **Group Policy** design, LDAPS, and DFS — plus internal and external **DNS**, DHCP, and all inbound and outbound SMTP routing.",
    "Own the storage estate across five platforms — Nimble, NetApp, EMC, Compellent, and Isilon — including capacity planning and performance analysis.",
    "**Converted the WAN from MPLS to SD-WAN**, running FortiGate over Cisco infrastructure and replacing fixed carrier circuits with a design that routes around failures on its own.",
    "Automated certificate issuance and renewal across three platforms — Windows and Active Directory, F5 BIG-IP, and Oracle Cloud Infrastructure — closing off the expiry outages that come with tracking certificates by hand.",
    "Deployed **Windows Event Forwarding** for centralized event collection and built account lockout analytics, making a recurring support burden diagnosable at the source instead of one reset ticket at a time.",
    "Automated server and printer deployment in PowerShell, and replaced a hardware inventory that had previously been collected by hand.",
    "Write the **as-built documentation, network and process diagrams, and backup and recovery runbooks** the team relies on, and implement all changes under formal change control.",
    "Third-level escalation point for complex infrastructure issues across compute, storage, networking, and hardware; after-hours on-call on a six-week rotation.",
]:
    bullet(d, b)

job(d, "Skype for Business / Teams Voice Engineer — Smart Shared Services",
    "May 2018 – April 2019")
for b in [
    "Delivered **client engagements** as the voice engineering lead — ran readiness and telephony assessments, produced the designs, and executed the migrations.",
    "Migrated enterprise voice from Lync to Skype for Business, and configured Sonus session border controllers for SIP trunking.",
    "Automated provisioning and policy enforcement in **PowerShell**, replacing per-user manual configuration.",
    "Translated business requirements into technical designs, then handed clients the operational documentation and **knowledge transfer** needed to run them.",
]:
    bullet(d, b)

job(d, "IT Director — Village Communities", "Westerville, OH  |  July 2011 – August 2017")
for b in [
    "Hired to run IT for the organization and owned it outright — infrastructure, budget, vendor contracts, and services for 150–250 employees and roughly 500 network accounts.",
    "**Consolidated physical servers onto Hyper-V**, and added a Nimble SAN for VM storage with site-to-site replication and backup.",
    "Led the JD Edwards EnterpriseOne conversion from 8.11 to 9.1 with the CFO, carrying the full CNC workload — package builds, OCM mappings, object promotions, ESUs and tools releases, environment creation and refresh, and Server Manager administration.",
    "Upgraded a multi-domain Windows environment across Server 2000, 2003, and 2012, and restructured Exchange by department and company.",
    "Deployed **Peplink SpeedFusion** bonded VPN across sites for circuit failover and throughput — an SD-WAN approach before the term was in common use — over Cisco WAN routing with OSPF, EIGRP, and RIP.",
    "Re-architected the network into segmented layers, eliminating broadcast storms and separating management traffic from end users.",
    "Operated a centralized ISP service for other tenants in the buildings, making IT a revenue line rather than only a cost center.",
]:
    bullet(d, b)

job(d, "Systems Administrator — ADiO Pharmacy", "October 2009 – July 2011")
for b in [
    "Managed Exchange across multiple domains, plus SharePoint, web, terminal, and domain servers.",
    "Deployed VMware Server and installed Unix servers inside Windows 2003 and 2008 domains with **failover clustering**; owned backups across the estate.",
    "Scripted routine administration in VBScript and PowerShell, and administered VPN, wireless, NAS storage, and SonicWall firewalls.",
]:
    bullet(d, b)

job(d, "Earlier", "")
body(d, "**Network Administrator**, Peer Assist (January – October 2009) — network and "
        "security administration, VOIP, cabling, and SQL databases behind in-house "
        "applications.", after=3)
body(d, "**Customer Service Representative**, Definitive Homecare Solutions (October 2007 "
        "– November 2008) — supported the CPR+ client-server application and built SQL "
        "reports against its data model.", after=0)

section(d, "Technical Skills")
for b in [
    "**Virtualization & HCI:** Hyper-V, Windows Failover Clustering, Cluster Shared Volumes, live migration, Storage Spaces Direct, SCVMM, Windows Admin Center, VMware vSphere/vCenter/ESXi/vSAN, Nutanix, Citrix XenApp, VDI",
    "**Windows Server:** Server 2000 through current, Active Directory, Group Policy, DNS, DHCP, LDAPS, DFS, certificate services, SMTP",
    "**Cloud & Hybrid:** Azure IaaS, Oracle Cloud Infrastructure, Microsoft 365, Exchange, SharePoint, Teams",
    "**Storage:** Storage Spaces Direct, vSAN, Nimble, NetApp, EMC, Compellent, Isilon",
    "**Automation:** PowerShell, VBScript; AI-assisted scripting and design work (Claude, ChatGPT, Gemini, Google AI Studio)",
    "**Networking:** Cisco, FortiGate SD-WAN, Peplink SpeedFusion, MPLS, F5 BIG-IP, OSPF, EIGRP, RIP, VLAN segmentation, spanning tree, VPN",
    "**Monitoring & Operations:** Windows Event Forwarding, System Center (Operations Manager, Data Protection Manager, Service Manager), change control, as-built documentation, runbooks",
    "**Voice:** Skype for Business, Teams Voice, Sonus SBC, SIP, H.323, Avaya",
    "**Business Systems:** JD Edwards EnterpriseOne (CNC), SQL Server, Yardi, PeopleSoft",
]:
    bullet(d, b)

section(d, "Education")
body(d, "**Columbus State Community College** — Computer Science coursework: C, COBOL, "
        "Assembler, JCL, Visual Basic, AS/400 CL", after=0)

d.save(f"{OUT}/Michael-George-Resume.docx")

# --------------------------------------------------------- cover letter

c = new_doc(0.75)
c.styles["Normal"].font.size = Pt(10.5)
name(c, "MICHAEL GEORGE")
contact(c, "Columbus, Ohio  |  mdgmikegeorgej@gmail.com  |  330-313-5604  |  "
           "linkedin.com/in/michael-george-02866411", rule=True)

body(c, "**Re: Senior Technical Consultant, Modern Data Center — Azure Local & Hyper-V**",
     after=11)
body(c, "Dear AHEAD Hiring Team,", after=10)
for para in [
    "I led my company's migration off VMware onto Hyper-V, and I run hyper-converged hosts on Storage Spaces Direct today. That covers a good share of what this role is asking for, so I'll be direct about both what I bring and what I don't.",
    "At Overhead Door Corporation I manage 500+ Windows and Linux servers across Hyper-V, VMware, and Nutanix, spread over multiple physical sites. I planned and executed the VMware-to-Hyper-V migration without interrupting production. I built the hyper-converged Storage Spaces Direct hosts and the failover clusters behind our high-availability workloads, standardized day-to-day management on Windows Admin Center, and automate the estate in PowerShell — including certificate issuance and renewal across Windows, F5 BIG-IP, and Oracle Cloud Infrastructure. I own Active Directory, DNS, and Group Policy end to end, and I converted our WAN from MPLS to SD-WAN.",
    "The gap, stated plainly: I have not deployed Azure Local itself. I have the substrate — Storage Spaces Direct, failover clustering, Windows Admin Center, Azure IaaS — but not the Arc registration, Network ATC, and cluster validation layer on top of it. That's a defined body of work rather than a new discipline, and it's the first thing I'd close. I'd rather say so here than have it surface in a technical screen.",
    "I've also worked this from the delivery side. At Smart Shared Services I ran client engagements as the voice engineering lead — readiness assessments, designs, migrations, and the documentation and knowledge transfer clients needed to operate what I built. Before that I spent six years as IT Director at Village Communities, owning the budget and vendor relationships and running a JD Edwards EnterpriseOne conversion alongside the CFO. Consulting delivery and talking to executives about risk and dependencies aren't new territory.",
    "One last thing, because your posting mentions it in three separate places: documentation is the part of this work I'm known for. Runbooks, as-builts, and migration procedures are usually the first thing to get skipped under schedule pressure, and they're the reason a handoff either holds or doesn't.",
    "I'd welcome the conversation.",
]:
    body(c, para, after=10)
body(c, "Michael George", after=0)

c.save(f"{OUT}/Michael-George-Cover-Letter-AHEAD.docx")
print("built both docx")
